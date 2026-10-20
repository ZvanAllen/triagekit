"""Desktop review tool (tkinter) for a triage coordinator.

Workflow:
  1. Import a referral batch (CSV).
  2. Set the batch waitlist status (system state shared by every referral).
  3. Run the deterministic engine.
  4. Review results in a sortable / filterable table; inspect any case.
  5. Optionally override an individual decision (with a note).
  6. Export results (CSV, Excel, or a printable summary).

Launch with ``triagekit-app`` (after ``pip install "triagekit[app]"``) or
``python -m triagekit.app``. tkinter ships with CPython; ``openpyxl`` is only
needed for the Excel export.

The triage logic lives in :func:`triage_referrals` (no GUI dependency) so it can
be tested headlessly; the ``TriageApp`` class is a thin Tk shell over it.
"""

from __future__ import annotations

import csv
from datetime import datetime
from importlib import resources
from pathlib import Path

from .engine import TriageEngine
from .factory import build_clinic_engine

APP_TITLE = "triagekit — Referral Triage Review"

DECISION_OPTIONS = ["ACCEPT", "DECLINE", "CONTACT_REQUIRED", "REFER", "REDIRECT_YOUTH"]
DECISION_COLORS = {
    "ACCEPT": "#2E7D32",
    "DECLINE": "#C62828",
    "CONTACT_REQUIRED": "#EF6C00",
    "REFER": "#6A1B9A",
    "REDIRECT_YOUTH": "#1565C0",
}
# Spreadsheet fills for the Excel export (lighter shades of the above).
DECISION_FILLS = {
    "ACCEPT": "C8E6C9",
    "DECLINE": "FFCDD2",
    "CONTACT_REQUIRED": "FFE0B2",
    "REFER": "E1BEE7",
    "REDIRECT_YOUTH": "BBDEFB",
}

# UI label -> referral field the batch waitlist toggle writes.
WAITLIST_FIELDS = [
    ("Medication open", "waitlist.medication_open"),
    ("Procedure open", "waitlist.procedure_open"),
    ("Counselling open", "waitlist.counselling_open"),
    ("Peer support open", "waitlist.peer_open"),
]

# Referral fields shown in the case-detail panel (label, key).
_DETAIL_FIELDS = [
    ("Age", "age"),
    ("Postal FSA", "postal_fsa"),
    ("Out-of-region insurance", "out_of_region_insurance"),
    ("In catchment", "in_catchment"),
    ("Connected PCP", "connected_pcp"),
    ("Virtual possible", "virtual_possible"),
    ("Can attend partner site", "can_attend_partner_site"),
    ("Prior treatment", "prior_treatment"),
    ("Comorbidity uncontrolled", "comorbidity_uncontrolled"),
    ("Procedure type", "procedure_type_specified"),
]
_REQUEST_FIELDS = [
    ("Med. initiation", "request.medication_initiation"),
    ("Med. maintenance", "request.medication_maintenance"),
    ("Procedure", "request.procedure"),
    ("Counselling", "request.counselling"),
    ("Peer support", "request.peer_support"),
]


# ---------------------------------------------------------------------------
# GUI-free logic (unit-testable)
# ---------------------------------------------------------------------------

def load_referrals_csv(path: str | Path) -> list[dict]:
    with open(path, encoding="utf-8", newline="") as f:
        return list(csv.DictReader(f))


def sample_referrals() -> list[dict]:
    """Load the small synthetic sample bundled with the package."""
    sample = Path(resources.files("triagekit") / "data" / "sample_referrals.csv")
    return load_referrals_csv(sample)


def triage_referrals(
    referrals: list[dict],
    waitlists: dict,
    engine: TriageEngine | None = None,
) -> list[dict]:
    """Run the engine over a batch, applying shared waitlist state to each row.

    Returns display-ready result dicts (engine output + row/name/age/id plus
    empty override slots and the source referral under ``_referral``). Pure
    function — no Tk, no I/O beyond the engine.
    """
    engine = engine or build_clinic_engine()
    results = []
    for i, raw in enumerate(referrals):
        ref = dict(raw)
        ref.update(waitlists)
        out = engine.run(ref).to_dict()
        out["row"] = i + 1
        out["full_name"] = ref.get("full_name") or f"Row {i + 1}"
        out["age"] = ref.get("age", "")
        out["referral_id"] = ref.get("referral_id", "")
        out["override_decision"] = ""
        out["override_notes"] = ""
        out["_referral"] = ref
        results.append(out)
    return results


def final_decision(result: dict) -> str:
    return result.get("override_decision") or result["decision"]


def export_rows(results: list[dict]) -> list[dict]:
    """Flatten results into serializable rows (lists joined to strings)."""
    rows = []
    for r in results:
        rows.append({
            "row": r["row"],
            "referral_id": r.get("referral_id", ""),
            "full_name": r["full_name"],
            "age": r.get("age", ""),
            "decision": r["decision"],
            "override_decision": r.get("override_decision", ""),
            "final_decision": final_decision(r),
            "accepted_services": ", ".join(r["accepted_services"]),
            "declined_services": ", ".join(r["declined_services"]),
            "refer_out_services": ", ".join(r["refer_out_services"]),
            "follow_up_required": "Yes" if r["follow_up_required"] else "No",
            "contradictions": "; ".join(r.get("contradictions", [])),
            "notes": r["notes"],
            "override_notes": r.get("override_notes", ""),
        })
    return rows


EXPORT_FIELDS = [
    "row", "referral_id", "full_name", "age",
    "decision", "override_decision", "final_decision",
    "accepted_services", "declined_services", "refer_out_services",
    "follow_up_required", "contradictions", "notes", "override_notes",
]


def write_csv(results: list[dict], path: str | Path) -> None:
    with open(path, "w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=EXPORT_FIELDS, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(export_rows(results))


def printable_summary(results: list[dict], waitlists: dict, source: str = "") -> str:
    from collections import Counter

    lines = ["Referral Triage Results"]
    lines.append(f"Generated: {datetime.now().strftime('%B %d, %Y at %I:%M %p')}")
    if source:
        lines.append(f"Source: {source}")
    lines.append(f"Cases: {len(results)}")
    wl = ", ".join(f"{label.split(' open')[0]}={'Open' if waitlists.get(key) else 'Closed'}"
                   for label, key in WAITLIST_FIELDS)
    lines.append(f"Waitlist: {wl}")
    lines.append("=" * 72)

    dist = Counter(final_decision(r) for r in results)
    lines.append("\nSUMMARY")
    for d in DECISION_OPTIONS:
        if dist.get(d):
            lines.append(f"  {d:<22} {dist[d]}")
    lines.append(f"  {'TOTAL':<22} {len(results)}")
    lines.append("\n" + "=" * 72)

    for r in results:
        decided = final_decision(r)
        lines.append("")
        lines.append(f"Case #{r['row']}: {r['full_name']}  (Age {r.get('age', '?')})")
        lines.append(f"  Decision: {decided}" + (" [OVERRIDDEN]" if r.get("override_decision") else ""))
        for label, key in (("Accepted", "accepted_services"),
                           ("Declined", "declined_services"),
                           ("Referred", "refer_out_services")):
            if r[key]:
                lines.append(f"  {label}: {', '.join(r[key])}")
        if r["follow_up_required"]:
            lines.append("  Follow-up required: Yes")
        if r.get("contradictions"):
            lines.append(f"  Contradictions: {'; '.join(r['contradictions'])}")
        lines.append(f"  Notes: {r['notes']}")
        if r.get("override_notes"):
            lines.append(f"  Override notes: {r['override_notes']}")
        lines.append("  " + "-" * 60)
    return "\n".join(lines)


# ---------------------------------------------------------------------------
# Tk application
# ---------------------------------------------------------------------------

def _build_app():
    """Import Tk lazily and return a constructed app instance."""
    import tkinter as tk
    from tkinter import filedialog, messagebox, scrolledtext, ttk

    class TriageApp(tk.Tk):
        def __init__(self):
            super().__init__()
            self.title(APP_TITLE)
            self.geometry("1280x800")
            self.minsize(900, 600)

            self.engine = build_clinic_engine()
            self.referrals: list[dict] = []
            self.results: list[dict] = []
            self.current_filter = "ALL"
            self.source_name = ""
            self.selected_row: int | None = None

            self.wl_vars = {key: tk.BooleanVar(value=True) for _, key in WAITLIST_FIELDS}

            self._ttk = ttk
            self._tk = tk
            self._dialogs = (filedialog, messagebox)
            self._scrolledtext = scrolledtext
            self._build_ui()

        # -- UI ---------------------------------------------------------
        def _build_ui(self):
            tk, ttk = self._tk, self._ttk

            style = ttk.Style(self)
            style.configure("Treeview", rowheight=26)
            style.configure("Treeview.Heading", font=("", 9, "bold"))

            menubar = tk.Menu(self)
            file_menu = tk.Menu(menubar, tearoff=0)
            file_menu.add_command(label="Import CSV...", command=self._import_csv, accelerator="Ctrl+O")
            file_menu.add_command(label="Load sample data", command=self._load_sample)
            file_menu.add_separator()
            file_menu.add_command(label="Export CSV...", command=self._export_csv)
            file_menu.add_command(label="Export Excel...", command=self._export_excel)
            file_menu.add_command(label="Export printable summary...", command=self._export_printable)
            file_menu.add_separator()
            file_menu.add_command(label="Exit", command=self.destroy)
            menubar.add_cascade(label="File", menu=file_menu)
            self.config(menu=menubar)
            self.bind_all("<Control-o>", lambda e: self._import_csv())

            top = ttk.LabelFrame(self, text="Batch settings", padding=10)
            top.pack(fill="x", padx=10, pady=(10, 5))

            wl_frame = ttk.Frame(top)
            wl_frame.pack(side="left", padx=(0, 30))
            ttk.Label(wl_frame, text="Waitlist status:", font=("", 10, "bold")).grid(
                row=0, column=0, columnspan=4, sticky="w", pady=(0, 5))
            for c, (label, key) in enumerate(WAITLIST_FIELDS):
                ttk.Checkbutton(wl_frame, text=label, variable=self.wl_vars[key]).grid(
                    row=1, column=c, padx=5, sticky="w")

            btns = ttk.Frame(top)
            btns.pack(side="right")
            ttk.Button(btns, text="Import CSV", command=self._import_csv).pack(side="left", padx=5)
            self.run_btn = ttk.Button(btns, text="Run triage", command=self._run, state="disabled")
            self.run_btn.pack(side="left", padx=5)
            self.status = ttk.Label(top, text="No data loaded", foreground="gray")
            self.status.pack(side="left", padx=20)

            filt = ttk.Frame(self, padding=(10, 5))
            filt.pack(fill="x")
            ttk.Label(filt, text="Filter:", font=("", 9, "bold")).pack(side="left")
            for d in ["ALL"] + DECISION_OPTIONS:
                ttk.Button(filt, text=d, command=lambda d=d: self._filter(d)).pack(side="left", padx=2)
            self.count = ttk.Label(filt, text="")
            self.count.pack(side="right")

            paned = ttk.PanedWindow(self, orient="vertical")
            paned.pack(fill="both", expand=True, padx=10, pady=5)

            table = ttk.Frame(paned)
            paned.add(table, weight=3)
            cols = ("row", "name", "age", "decision", "accepted", "declined", "referred", "fu", "notes")
            self.tree = ttk.Treeview(table, columns=cols, show="headings", selectmode="browse")
            for cid, (head, w) in {
                "row": ("#", 40), "name": ("Name", 150), "age": ("Age", 45),
                "decision": ("Decision", 150), "accepted": ("Accepted", 190),
                "declined": ("Declined", 190), "referred": ("Referred", 160),
                "fu": ("F/U", 45), "notes": ("Notes", 340),
            }.items():
                self.tree.heading(cid, text=head, command=lambda c=cid: self._sort(c))
                self.tree.column(cid, width=w, minwidth=30)
            vsb = ttk.Scrollbar(table, orient="vertical", command=self.tree.yview)
            hsb = ttk.Scrollbar(table, orient="horizontal", command=self.tree.xview)
            self.tree.configure(yscrollcommand=vsb.set, xscrollcommand=hsb.set)
            self.tree.grid(row=0, column=0, sticky="nsew")
            vsb.grid(row=0, column=1, sticky="ns")
            hsb.grid(row=1, column=0, sticky="ew")
            table.grid_rowconfigure(0, weight=1)
            table.grid_columnconfigure(0, weight=1)
            self.tree.bind("<<TreeviewSelect>>", self._on_select)
            for d, color in DECISION_COLORS.items():
                self.tree.tag_configure(d, foreground=color)

            detail = ttk.LabelFrame(paned, text="Case detail", padding=10)
            paned.add(detail, weight=2)
            self.detail = self._scrolledtext.ScrolledText(detail, wrap="word", height=8, font=("Consolas", 10))
            self.detail.pack(side="left", fill="both", expand=True)

            ov = ttk.LabelFrame(detail, text="Override", padding=10)
            ov.pack(side="right", fill="y", padx=(10, 0))
            ttk.Label(ov, text="Decision:").pack(anchor="w")
            self.ov_decision = ttk.Combobox(ov, values=DECISION_OPTIONS, state="readonly", width=20)
            self.ov_decision.pack(pady=(0, 5))
            ttk.Label(ov, text="Note:").pack(anchor="w")
            self.ov_notes = tk.Text(ov, width=26, height=4, font=("", 9))
            self.ov_notes.pack(pady=(0, 5))
            self.ov_btn = ttk.Button(ov, text="Apply override", command=self._apply_override, state="disabled")
            self.ov_btn.pack(anchor="w")
            ttk.Label(ov, text="Overrides are flagged with * and\nrecorded in every export.",
                      foreground="gray", font=("", 8)).pack(anchor="w", pady=(8, 0))

        # -- data in ----------------------------------------------------
        def _import_csv(self):
            filedialog, messagebox = self._dialogs
            path = filedialog.askopenfilename(
                title="Import referral CSV", filetypes=[("CSV files", "*.csv"), ("All files", "*.*")])
            if not path:
                return
            try:
                self.referrals = load_referrals_csv(path)
            except Exception as e:  # noqa: BLE001
                messagebox.showerror("Import error", f"Failed to read CSV:\n{e}")
                return
            self.source_name = Path(path).name
            self._after_load()

        def _load_sample(self):
            try:
                self.referrals = sample_referrals()
                self.source_name = "sample_referrals.csv"
                self._after_load()
            except Exception as e:  # noqa: BLE001
                self._dialogs[1].showerror("Sample error", str(e))

        def _after_load(self):
            self.results = []
            self._clear_table()
            self.status.config(text=f"Loaded {len(self.referrals)} referrals from {self.source_name}",
                               foreground="black")
            self.run_btn.config(state="normal")
            self.count.config(text="")

        # -- run --------------------------------------------------------
        def _run(self):
            if not self.referrals:
                self._dialogs[1].showwarning("No data", "Import a CSV first.")
                return
            waitlists = {key: var.get() for key, var in self.wl_vars.items()}
            self.results = triage_referrals(self.referrals, waitlists, self.engine)
            self._populate()
            from collections import Counter
            dist = Counter(final_decision(r) for r in self.results)
            summary = ", ".join(f"{d}: {n}" for d, n in sorted(dist.items()))
            self.status.config(text=f"Triage complete — {len(self.results)} cases. {summary}",
                               foreground=DECISION_COLORS["ACCEPT"])

        # -- table ------------------------------------------------------
        def _clear_table(self):
            for item in self.tree.get_children():
                self.tree.delete(item)
            self.detail.delete("1.0", "end")

        def _populate(self):
            self._clear_table()
            for r in self.results:
                decided = final_decision(r)
                if self.current_filter != "ALL" and decided != self.current_filter:
                    continue
                self.tree.insert("", "end", iid=str(r["row"]), tags=(decided,), values=(
                    r["row"], r["full_name"], r["age"],
                    decided + (" *" if r.get("override_decision") else ""),
                    ", ".join(r["accepted_services"]),
                    ", ".join(r["declined_services"]),
                    ", ".join(r["refer_out_services"]),
                    "Yes" if r["follow_up_required"] else "",
                    r["notes"][:100],
                ))
            self.count.config(text=f"Showing {len(self.tree.get_children())}/{len(self.results)}")

        def _filter(self, decision):
            self.current_filter = decision
            if self.results:
                self._populate()

        def _sort(self, col):
            items = [(self.tree.set(k, col), k) for k in self.tree.get_children()]
            try:
                items.sort(key=lambda t: int(t[0]))
            except ValueError:
                items.sort(key=lambda t: t[0].lower())
            for i, (_, k) in enumerate(items):
                self.tree.move(k, "", i)

        # -- detail / override ------------------------------------------
        def _on_select(self, _event):
            sel = self.tree.selection()
            if not sel:
                return
            row = int(sel[0])
            result = next((r for r in self.results if r["row"] == row), None)
            if not result:
                return
            self.selected_row = row
            self.ov_btn.config(state="normal")
            self.detail.delete("1.0", "end")
            self.detail.insert("1.0", self._detail_text(result))
            self.ov_decision.set(final_decision(result))
            self.ov_notes.delete("1.0", "end")
            if result.get("override_notes"):
                self.ov_notes.insert("1.0", result["override_notes"])

        def _detail_text(self, result) -> str:
            ref = result.get("_referral", {})
            decided = final_decision(result)
            out = [f"{'=' * 60}", f"  CASE #{result['row']}: {result['full_name']}", "=" * 60, ""]
            out.append(f"  DECISION:  {decided}")
            if result.get("override_decision"):
                out.append(f"  (Engine: {result['decision']}, overridden)")
            out.append("")
            for label, key in (("Accepted", "accepted_services"),
                               ("Declined", "declined_services"),
                               ("Referred", "refer_out_services")):
                if result[key]:
                    out.append(f"  {label}:  {', '.join(result[key])}")
            if result["follow_up_required"]:
                out.append("  FOLLOW-UP REQUIRED: Yes")
            if result.get("contradictions"):
                out.append("  CONTRADICTIONS:")
                out.extend(f"    - {c}" for c in result["contradictions"])
            out.append("")
            if result["referral_destination"]:
                out.append(f"  DESTINATION: {result['referral_destination']}")
            if result["routing_instruction"]:
                out.append(f"  ROUTING: {result['routing_instruction']}")
            out.append(f"  NOTES: {result['notes']}")
            if result.get("override_notes"):
                out.append(f"  OVERRIDE NOTES: {result['override_notes']}")
            out += ["", "  " + "-" * 58, "  REFERRAL DATA:"]
            for label, key in _DETAIL_FIELDS:
                out.append(f"    {label}: {ref.get(key, '')}")
            reqs = ", ".join(f"{label}={ref.get(key, '')}" for label, key in _REQUEST_FIELDS)
            out.append(f"    Requests: {reqs}")
            out.append(f"    Meds: {ref.get('meds_text', '')}")
            out.append(f"    History: {ref.get('med_history_text', '')}")
            if ref.get("referral_id"):
                out.append(f"    Referral ID: {ref['referral_id']}")
            return "\n".join(out)

        def _apply_override(self):
            if self.selected_row is None:
                return
            result = next((r for r in self.results if r["row"] == self.selected_row), None)
            if not result:
                return
            new_decision = self.ov_decision.get()
            result["override_decision"] = new_decision if new_decision != result["decision"] else ""
            result["override_notes"] = self.ov_notes.get("1.0", "end").strip()
            self._populate()
            try:
                self.tree.selection_set(str(self.selected_row))
                self._on_select(None)
            except self._tk.TclError:
                pass

        # -- export -----------------------------------------------------
        def _ask_save(self, ext, filetypes, prefix):
            return self._dialogs[0].asksaveasfilename(
                title="Export", defaultextension=ext, filetypes=filetypes,
                initialfile=f"{prefix}_{datetime.now():%Y%m%d_%H%M}{ext}")

        def _need_results(self) -> bool:
            if not self.results:
                self._dialogs[1].showwarning("No results", "Run triage first.")
                return False
            return True

        def _export_csv(self):
            if not self._need_results():
                return
            path = self._ask_save(".csv", [("CSV files", "*.csv")], "triage_results")
            if path:
                write_csv(self.results, path)
                self._dialogs[1].showinfo("Exported", f"Results saved to:\n{path}")

        def _export_excel(self):
            if not self._need_results():
                return
            try:
                import openpyxl
                from openpyxl.styles import Font, PatternFill
            except ImportError:
                self._dialogs[1].showerror("Missing package", 'Install openpyxl:\n  pip install "triagekit[app]"')
                return
            path = self._ask_save(".xlsx", [("Excel files", "*.xlsx")], "triage_results")
            if not path:
                return
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Triage results"
            rows = export_rows(self.results)
            headers = list(EXPORT_FIELDS)
            for c, h in enumerate(headers, 1):
                cell = ws.cell(row=1, column=c, value=h)
                cell.font = Font(bold=True)
                cell.fill = PatternFill("solid", fgColor="D9E1F2")
            final_col = headers.index("final_decision") + 1
            for i, row in enumerate(rows, 2):
                for c, h in enumerate(headers, 1):
                    ws.cell(row=i, column=c, value=row[h])
                fill = DECISION_FILLS.get(row["final_decision"])
                if fill:
                    ws.cell(row=i, column=final_col).fill = PatternFill("solid", fgColor=fill)
            ws.freeze_panes = "A2"
            wb.save(path)
            self._dialogs[1].showinfo("Exported", f"Results saved to:\n{path}")

        def _export_printable(self):
            if not self._need_results():
                return
            path = self._ask_save(".txt", [("Text files", "*.txt")], "triage_summary")
            if not path:
                return
            waitlists = {key: var.get() for key, var in self.wl_vars.items()}
            Path(path).write_text(
                printable_summary(self.results, waitlists, self.source_name), encoding="utf-8")
            self._dialogs[1].showinfo("Exported", f"Summary saved to:\n{path}")

    return TriageApp()


def main() -> None:
    _build_app().mainloop()


if __name__ == "__main__":
    main()
